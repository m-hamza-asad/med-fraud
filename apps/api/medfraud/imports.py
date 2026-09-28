from __future__ import annotations

import csv
import hashlib
import io
import json
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable

from sqlalchemy import select
from sqlalchemy.orm import Session

from .models import CanonicalFact, Claim, ImportBatch, ImportPreview

from openpyxl import load_workbook


PROFILES = {"canonical", "shafafiya", "eclaimlink", "synthetic_uae"}
REQUIRED = {"source_claim_id", "member_token", "provider_token", "claim_type", "service_date", "submitted_amount", "net_amount"}
BLOCKED_LOGIC_FIELDS = {"signals", "fraud_label", "fraud_confidence", "synthetic_score"}
SYNTHETIC_LABEL_FIELDS = {"fraud_label", "fraud_type", "fraud_confidence", "ground_truth_source"}
SYNTHETIC_REQUIRED = {
    "claim_id", "patient_id", "hospital_id", "date_of_admission", "date_of_discharge", "date_of_claim",
    "claim_amount_requested_aed", "claim_amount_approved_aed",
}


def checksum(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def safe_filename(name: str) -> str:
    return Path(name).name.replace("\x00", "")[:255]


def _csv_rows(data: bytes) -> list[dict[str, Any]]:
    text = data.decode("utf-8-sig")
    return list(csv.DictReader(io.StringIO(text)))


def _xlsx_rows(data: bytes) -> list[dict[str, Any]]:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    if "claim_header" not in workbook.sheetnames:
        raise ValueError("Missing required sheet: claim_header")
    sheet = workbook["claim_header"]
    headers = [str(v) if v is not None else "" for v in next(sheet.iter_rows(values_only=True))]
    return [dict(zip(headers, row, strict=False)) for row in sheet.iter_rows(values_only=True)]


def parse_rows(filename: str, data: bytes) -> list[dict[str, Any]]:
    lower = filename.lower()
    if lower.endswith(".csv"):
        return _csv_rows(data)
    if lower.endswith(".xlsx"):
        return _xlsx_rows(data)
    raise ValueError("Only .csv and .xlsx files are accepted")


def _boolean(value: Any) -> bool | None:
    text = str(value).strip().casefold()
    if text in {"true", "1", "yes"}: return True
    if text in {"false", "0", "no"}: return False
    return None


def _day_first(value: Any) -> date:
    if isinstance(value, datetime): return value.date()
    if isinstance(value, date): return value
    text = str(value).strip()
    for pattern in ("%d/%m/%Y", "%Y-%m-%d"):
        try: return datetime.strptime(text, pattern).date()
        except ValueError: pass
    raise ValueError(f"invalid date '{text}' (expected DD/MM/YYYY or YYYY-MM-DD)")


def map_synthetic_uae_rows(rows: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """Map the explicit synthetic claim-header schema; labels remain held out."""
    if not rows: return [], []
    missing = sorted(SYNTHETIC_REQUIRED - set(rows[0]))
    if missing:
        return [], [{"severity": "error", "row_number": 1, "column_name": column,
                     "code": "MISSING_SYNTHETIC_COLUMN", "message": f"Required synthetic column '{column}' is missing",
                     "guidance": "Use the synthetic_uae profile after preparing the AED-labelled fixture."} for column in missing]
    mapped: list[dict[str, Any]] = []
    issues: list[dict[str, Any]] = []
    for number, raw in enumerate(rows, start=2):
        try:
            admission = _day_first(raw["date_of_admission"]); discharge = _day_first(raw["date_of_discharge"])
            claim_date = _day_first(raw["date_of_claim"]); policy_days = int(raw.get("days_since_policy_start") or 0)
            requested = Decimal(str(raw["claim_amount_requested_aed"])); approved = Decimal(str(raw["claim_amount_approved_aed"])); tpa = str(raw.get("tpa") or "UNSPECIFIED_TPA").strip()
            facts = {
                "available_datasets": ["claim_header", "encounter", "diagnosis", "coverage_period", "provider", "provider_status", "claim_history", "observation", "provider_features", "provider_network", "policy_enrollment"],
                "claim": {"submission_date": claim_date.isoformat(), "gross_amount": str(requested), "approved_amount": str(approved), "requested_payable": str(requested)},
                "encounter": {"admission_date": admission.isoformat(), "discharge_date": discharge.isoformat(), "length_of_stay_days": int(raw.get("length_of_stay_days") or 0)},
                "diagnosis": {"primary_code": str(raw.get("diagnosis_primary") or ""), "description": str(raw.get("diagnosis_description") or ""), "code_matches_procedure": _boolean(raw.get("icd_code_matches_procedure"))},
                "coverage": {"policy_type": str(raw.get("policy_type") or ""), "valid_from": (claim_date - timedelta(days=policy_days)).isoformat(), "valid_to": "2099-12-31", "network_id": tpa},
                "provider": {"blacklist_flag": _boolean(raw.get("provider_blacklist_flag")), "network_id": tpa},
                "history": {"readmission_gap_days": int(raw.get("discharge_readmit_gap_days") or 0), "same_event_insurer_count": int(raw.get("num_insurers_same_event") or 0), "previous_reviewed_fraud_on_policy": _boolean(raw.get("previous_fraud_on_policy"))},
                "pharmacy": {"bill_ratio": str(raw.get("pharmacy_bill_ratio") or "")},
                "channel": {"agent_token": str(raw.get("agent_id") or ""), "tpa": tpa, "cashless": _boolean(raw.get("is_cashless"))},
                "mapping": {"profile": "synthetic_uae_v1", "currency": "AED", "labels_held_out": sorted(SYNTHETIC_LABEL_FIELDS)},
            }
            mapped.append({"source_claim_id": str(raw["claim_id"]).strip(), "member_token": str(raw["patient_id"]).strip(),
                           "provider_token": str(raw["hospital_id"]).strip(), "network_id": tpa,
                           "claim_type": "inpatient", "service_date": admission.isoformat(),
                           "submitted_amount": str(requested), "net_amount": str(approved), "paid_amount": "0",
                           "facts_json": json.dumps(facts)})
        except (ValueError, TypeError, InvalidOperation) as exc:
            issues.append({"severity": "error", "row_number": number, "column_name": None, "code": "SYNTHETIC_MAPPING_ERROR",
                           "message": str(exc), "guidance": "Correct the source row or mapping profile."})
    return mapped, issues


def validate_rows(rows: list[dict[str, Any]], filename: str, analysis_date: date) -> tuple[list[dict], list[dict]]:
    issues: list[dict] = []
    normalized: list[dict] = []
    if not rows:
        return [], [{"severity": "error", "row_number": None, "column_name": None, "code": "EMPTY_FILE", "message": "No claim rows were found", "guidance": "Add a header row and at least one claim."}]
    missing = REQUIRED - set(rows[0])
    for col in sorted(missing):
        issues.append({"severity": "error", "row_number": 1, "column_name": col, "code": "MISSING_COLUMN", "message": f"Required column '{col}' is missing", "guidance": "Use the matching downloadable template."})
    if missing:
        return [], issues
    seen: set[str] = set()
    cutoff = date(analysis_date.year - 5, analysis_date.month, min(analysis_date.day, 28))
    for offset, raw in enumerate(rows, start=2):
        try:
            service = raw["service_date"] if isinstance(raw["service_date"], date) else date.fromisoformat(str(raw["service_date"])[:10])
            submitted = Decimal(str(raw["submitted_amount"])).quantize(Decimal("0.01"))
            net = Decimal(str(raw["net_amount"])).quantize(Decimal("0.01"))
            paid = Decimal(str(raw.get("paid_amount") or 0)).quantize(Decimal("0.01"))
            source_id = str(raw["source_claim_id"]).strip()
            if source_id in seen:
                raise ValueError("duplicate source claim ID")
            seen.add(source_id)
            if service > analysis_date:
                raise ValueError("service date is after the analysis date")
            if service < cutoff:
                issues.append({"severity": "warning", "row_number": offset, "column_name": "service_date", "code": "OUTSIDE_WINDOW", "message": "Claim is older than the five-year analysis window", "guidance": "Keep for history only or choose a later analysis date."})
            facts_raw = raw.get("facts_json") or "{}"
            facts = facts_raw if isinstance(facts_raw, dict) else json.loads(str(facts_raw))
            if not isinstance(facts, dict):
                raise ValueError("facts_json must contain an object")
            if net > submitted:
                if facts.get("mapping", {}).get("profile") == "synthetic_uae_v1":
                    issues.append({"severity": "warning", "row_number": offset, "column_name": "net_amount",
                                   "code": "APPROVED_EXCEEDS_REQUESTED", "message": "Synthetic approved amount exceeds requested amount",
                                   "guidance": "Preserved as an anomaly for deterministic rule evaluation; verify against real source semantics."})
                else:
                    raise ValueError("net amount exceeds submitted amount")
            for field in BLOCKED_LOGIC_FIELDS:
                facts.pop(field, None)
            normalized.append({**raw, "source_claim_id": source_id, "service_date": service, "submitted_amount": submitted, "net_amount": net, "paid_amount": paid, "facts_json": json.dumps(facts), "valid_from": service})
        except (ValueError, TypeError, InvalidOperation, json.JSONDecodeError) as exc:
            issues.append({"severity": "error", "row_number": offset, "column_name": None, "code": "INVALID_ROW", "message": str(exc), "guidance": "Correct the row values and validate again."})
    for issue in issues:
        issue.update({"file": filename, "sheet": "claim_header" if filename.lower().endswith(".xlsx") else None})
    return normalized, issues


def preview_payload(rows: Iterable[dict[str, Any]]) -> str:
    """Serialize validated normalized rows without persisting operational claims."""
    return json.dumps(list(rows), default=lambda value: value.isoformat() if isinstance(value, (date, datetime)) else str(value))


def _fact_rows(claim: Claim, value: Any, path: str = "", dataset: str = "claim_header") -> list[CanonicalFact]:
    rows: list[CanonicalFact] = []
    if isinstance(value, dict):
        for key, child in value.items():
            if key in BLOCKED_LOGIC_FIELDS:
                continue
            child_path = f"{path}.{key}" if path else key
            child_dataset = child_path.split(".", 1)[0] if "." in child_path else dataset
            rows.extend(_fact_rows(claim, child, child_path, child_dataset))
        return rows
    if isinstance(value, list):
        value = json.dumps(value, default=str)
    kind = "boolean" if isinstance(value, bool) else "number" if isinstance(value, (int, float, Decimal)) else "date" if isinstance(value, date) else "text"
    kwargs: dict[str, Any] = {"value_type": kind, "value_text": None, "value_number": None, "value_date": None, "value_bool": None}
    if kind == "boolean": kwargs["value_bool"] = value
    elif kind == "number": kwargs["value_number"] = Decimal(str(value))
    elif kind == "date": kwargs["value_date"] = value
    else: kwargs["value_text"] = "" if value is None else str(value)
    rows.append(CanonicalFact(claim_id=claim.id, entity_type="claim", entity_key=claim.source_claim_id,
                              dataset=dataset, field_name=path, valid_from=claim.valid_from,
                              source_key=f"{claim.source_claim_id}:{path}", source=claim.source,
                              version_id=claim.version_id, **kwargs))
    return rows


def commit_preview(db: Session, batch: ImportBatch) -> int:
    """Atomically materialize one validated preview; safe to retry after commit."""
    if batch.status == "COMMITTED":
        return 0
    if batch.status != "VALIDATED":
        raise ValueError("Only a validated import can be committed")
    preview = db.scalar(select(ImportPreview).where(ImportPreview.batch_id == batch.id))
    if preview is None:
        raise ValueError("Validated import preview is unavailable")
    rows = json.loads(preview.normalized_rows_json)
    created = 0
    for start in range(0, len(rows), 500):
        pending: list[tuple[Claim, dict[str, Any]]] = []
        for raw in rows[start:start + 500]:
            facts = json.loads(raw.get("facts_json") or "{}")
            for field in BLOCKED_LOGIC_FIELDS:
                facts.pop(field, None)
            claim = Claim(source_profile=batch.profile, source_claim_id=str(raw["source_claim_id"]),
                          member_token=str(raw["member_token"]), provider_token=str(raw["provider_token"]),
                          network_id=str(raw.get("network_id") or "") or None, claim_type=str(raw["claim_type"]),
                          service_date=date.fromisoformat(str(raw["service_date"])[:10]),
                          submitted_amount=Decimal(str(raw["submitted_amount"])), net_amount=Decimal(str(raw["net_amount"])),
                          paid_amount=Decimal(str(raw.get("paid_amount") or 0)), facts_json=json.dumps(facts),
                          original_json=json.dumps({k: v for k, v in raw.items() if k not in BLOCKED_LOGIC_FIELDS}),
                          valid_from=date.fromisoformat(str(raw["valid_from"])[:10]), import_batch_id=batch.id)
            db.add(claim); pending.append((claim, facts))
        db.flush()
        for claim, facts in pending:
            canonical = {"claim": {"source_claim_id": claim.source_claim_id, "member_token": claim.member_token,
                                    "provider_token": claim.provider_token, "claim_type": claim.claim_type,
                                    "service_date": claim.service_date, "submitted_amount": claim.submitted_amount,
                                    "net_amount": claim.net_amount, "paid_amount": claim.paid_amount}, **facts}
            db.add_all(_fact_rows(claim, canonical))
        db.flush()
        created += len(pending)
    batch.status = "COMMITTED"
    return created

